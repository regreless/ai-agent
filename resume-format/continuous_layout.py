from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from pathlib import Path

base_path=Path(r'C:\AIApp\ai-agent\resume-format')
source_path=base_path/'刘谦-简历-标题同行版.docx'
output_path=base_path/'刘谦-简历-连续排版版.docx'
doc=Document(source_path)
for paragraph in list(doc.paragraphs):
 if not paragraph.text.strip() and not paragraph._p.xpath('.//w:drawing'):
  paragraph._element.getparent().remove(paragraph._element)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Paragraph']:
 paragraph_format=doc.styles[style_name].paragraph_format
 paragraph_format.space_before=Pt(0)
 paragraph_format.space_after=Pt(0)
 paragraph_format.keep_with_next=False
 paragraph_format.keep_together=False
 paragraph_format.widow_control=False
 paragraph_format.page_break_before=False
 paragraph_format.line_spacing=1.5
paragraphs=list(doc.paragraphs)
for paragraph_index,paragraph in enumerate(paragraphs):
 paragraph_format=paragraph.paragraph_format
 paragraph_format.space_before=Pt(1 if paragraph.style.name.startswith('Heading') else 0)
 paragraph_format.space_after=Pt(0)
 paragraph_format.keep_together=False
 paragraph_format.widow_control=False
 paragraph_format.page_break_before=False
 paragraph_format.line_spacing=1.5
 paragraph_format.keep_with_next=False
 if paragraph.style.name.startswith('Heading') and paragraph_index+1<len(paragraphs):
  paragraph_format.keep_with_next=not paragraphs[paragraph_index+1].style.name.startswith('Heading')
 for break_element in list(paragraph._p.xpath('.//w:br[@w:type="page"]')):
  break_element.getparent().remove(break_element)
assert ''.join(p.text for p in doc.paragraphs)==''.join(p.text for p in Document(source_path).paragraphs)
doc.save(output_path)
print('Saved continuous layout')
