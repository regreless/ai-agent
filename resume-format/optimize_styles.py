from pathlib import Path
import re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_PATH=Path(r'C:\AIApp\ai-agent\resume-format')
SOURCE_PATH=Path(r'C:\Users\Administrator\Desktop\刘谦-简历-连续排版版.docx')
OUTPUT_PATH=BASE_PATH/'刘谦-简历-样式优化版.docx'
BLUE='1F4E79'
BODY='202020'
GRAY='626B73'
source_doc=Document(SOURCE_PATH)
output_doc=Document()
section=output_doc.sections[0]
section.page_width=Cm(21)
section.page_height=Cm(29.7)
section.top_margin=Cm(1.45)
section.bottom_margin=Cm(1.45)
section.left_margin=Cm(1.5)
section.right_margin=Cm(1.5)
section.footer_distance=Cm(.65)
for grid in list(section._sectPr.findall(qn('w:docGrid'))):section._sectPr.remove(grid)

def set_font(target,size=10.5,color=BODY,bold=False,east_asia='等线'):
 target.font.name='Calibri'
 target.font.size=Pt(size)
 target.font.bold=bold
 target.font.color.rgb=RGBColor.from_string(color)
 properties=target._element.get_or_add_rPr()
 fonts=properties.get_or_add_rFonts()
 for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  fonts.attrib.pop(qn('w:'+key),None)
 fonts.set(qn('w:eastAsia'),east_asia)
 fonts.set(qn('w:ascii'),'Calibri')
 fonts.set(qn('w:hAnsi'),'Calibri')
 size_cs=properties.find(qn('w:szCs'))
 if size_cs is None:
  size_cs=OxmlElement('w:szCs')
  properties.append(size_cs)
 size_cs.set(qn('w:val'),str(int(size*2)))

for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Paragraph']:
 style=output_doc.styles[style_name]
 set_font(style)
 paragraph_format=style.paragraph_format
 paragraph_format.line_spacing=1.5
 paragraph_format.space_before=Pt(0)
 paragraph_format.space_after=Pt(0)
 paragraph_format.keep_together=False
 paragraph_format.keep_with_next=False
 paragraph_format.widow_control=True
for style_name,size,before in [('Heading 1',13,5),('Heading 2',11.5,3),('Heading 3',10.5,2)]:
 style=output_doc.styles[style_name]
 set_font(style,size,BLUE,True,'微软雅黑')
 style.paragraph_format.line_spacing=1.0
 style.paragraph_format.space_before=Pt(before)
 style.paragraph_format.space_after=Pt(1)
 style.paragraph_format.keep_with_next=True
set_font(output_doc.styles['Title'],22,'000000',True,'微软雅黑')
output_doc.styles['Title'].paragraph_format.line_spacing=1.0
output_doc.styles['Title'].paragraph_format.space_after=Pt(3)

text_records=[]
for paragraph_index,source_paragraph in enumerate(source_doc.paragraphs):
 text=source_paragraph.text.replace('**','').strip()
 if not text:continue
 if paragraph_index in [73,75,77] and not text.startswith('• '):text='• '+text
 if paragraph_index==77 and '交付结果：' in text:
  first_text,last_text=text.split('交付结果：',1)
  text_records.append((paragraph_index,'Normal',first_text))
  text_records.append((paragraph_index,'Normal','交付结果：'+last_text))
 else:text_records.append((paragraph_index,source_paragraph.style.name,text))

for paragraph_index,style_name,text in text_records:
 paragraph=output_doc.add_paragraph(style=style_name if style_name in ['Title','Heading 1','Heading 2','Heading 3'] else 'Normal')
 paragraph_format=paragraph.paragraph_format
 paragraph_format.page_break_before=False
 snap_element=OxmlElement('w:snapToGrid')
 snap_element.set(qn('w:val'),'0')
 paragraph._p.get_or_add_pPr().append(snap_element)
 if style_name=='Title':
  paragraph.add_run(text)
 elif paragraph_index==1:
  paragraph_format.line_spacing=1.0
  paragraph_format.space_after=Pt(3)
  set_font(paragraph.add_run(text),11.5,BLUE,True,'微软雅黑')
 elif paragraph_index==2:
  paragraph_format.line_spacing=1.0
  paragraph_format.space_after=Pt(4)
  set_font(paragraph.add_run(text),9.5,GRAY)
 elif style_name.startswith('Heading'):
  if '\t' in text:
   left_text,right_text=text.split('\t',1)
   paragraph_format.tab_stops.add_tab_stop(Cm(18.0),WD_TAB_ALIGNMENT.RIGHT)
   paragraph.add_run(left_text)
   set_font(paragraph.add_run('\t'+right_text),9.5,GRAY,False)
  else:paragraph.add_run(text)
 else:
  is_bullet=text.startswith('• ')
  if is_bullet:
   paragraph_format.left_indent=Cm(.32)
   paragraph_format.first_line_indent=Cm(-.32)
   set_font(paragraph.add_run('• '),10.5,GRAY)
   text=text[2:]
  match=re.match(r'^([^：]{1,35}：)(.*)$',text)
  if match:
   set_font(paragraph.add_run(match.group(1)),10.5,BLUE,True)
   set_font(paragraph.add_run(match.group(2)))
  else:set_font(paragraph.add_run(text))

paragraphs=output_doc.paragraphs
for paragraph_index,paragraph in enumerate(paragraphs[:-1]):
 if paragraph.style.name.startswith('Heading') and paragraphs[paragraph_index+1].style.name.startswith('Heading'):
  paragraph.paragraph_format.keep_with_next=True

footer_paragraph=section.footer.paragraphs[0]
footer_paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
field_run=OxmlElement('w:r')
field_properties=OxmlElement('w:rPr')
field_color=OxmlElement('w:color')
field_color.set(qn('w:val'),GRAY)
field_size=OxmlElement('w:sz')
field_size.set(qn('w:val'),'16')
field_properties.append(field_color)
field_properties.append(field_size)
field_run.append(field_properties)
field_text=OxmlElement('w:t')
field_text.text='1'
field_run.append(field_text)
field.append(field_run)
footer_paragraph._p.append(field)
for element in [output_doc.element,output_doc.styles.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
output_doc.core_properties.title='刘谦 简历'
output_doc.core_properties.author='刘谦'
output_doc.core_properties.subject=''
output_doc.save(OUTPUT_PATH)

def normalize_text(doc):
 return re.sub(r'[\s•*]','',''.join(p.text for p in doc.paragraphs))
assert normalize_text(source_doc)==normalize_text(output_doc),'Content changed'
print('Saved; content preserved')

