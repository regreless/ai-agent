from pathlib import Path
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_PATH=Path(r'C:\AIApp\ai-agent\resume-format')
SOURCE_PATH=BASE_PATH/'刘谦-简历-样式优化版.docx'
OUTPUT_PATH=BASE_PATH/'刘谦-简历-微软雅黑版.docx'
FONT_NAME='微软雅黑'
BLUE='1F4E79'
doc=Document(SOURCE_PATH)
in_projects=False
background_count=0
for paragraph in doc.paragraphs:
 if paragraph.text=='项目经验':in_projects=True
 if in_projects and paragraph.text=='项目背景':
  paragraph.clear()
  paragraph.style='Normal'
  paragraph.paragraph_format.keep_with_next=False
  paragraph.paragraph_format.line_spacing=1.5
  label_run=paragraph.add_run('项目背景：')
  label_run.bold=True
  label_run.font.color.rgb=RGBColor.from_string(BLUE)
  body_run=paragraph.add_run('面向急诊抢救、重症监护等临床协同场景，业务具有时间敏感、角色协作密集和数据准确性要求高的特点，需要将分散的诊疗数据组织为连续、可验证的业务流程，并保障高频迭代下的稳定性。')
  body_run.font.color.rgb=RGBColor.from_string('202020')
 if in_projects and paragraph.text.startswith('项目背景：'):background_count+=1

for part in doc.part.package.parts:
 element=getattr(part,'_element',None)
 if element is None:continue
 for fonts in element.iter(qn('w:rFonts')):
  for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   fonts.attrib.pop(qn('w:'+key),None)
  for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),FONT_NAME)
 for run_element in element.iter(qn('w:r')):
  properties=run_element.find(qn('w:rPr'))
  if properties is None:
   properties=OxmlElement('w:rPr')
   run_element.insert(0,properties)
  fonts=properties.find(qn('w:rFonts'))
  if fonts is None:
   fonts=OxmlElement('w:rFonts')
   properties.insert(0,fonts)
  for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),FONT_NAME)
for style in doc.styles:
 if hasattr(style,'font'):
  style.font.name=FONT_NAME
  fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
  for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
  for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),FONT_NAME)
assert background_count==2,background_count
assert all((p.paragraph_format.line_spacing or p.style.paragraph_format.line_spacing)==1.5 or p.style.name!='Normal' or i in [1,2] for i,p in enumerate(doc.paragraphs))
doc.save(OUTPUT_PATH)
print('Saved; project backgrounds:',background_count)

