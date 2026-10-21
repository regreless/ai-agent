from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
import re

output_path=r'C:\AIApp\ai-agent\resume-format\刘谦-简历-排版优化版.docx'
doc=Document(output_path)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Paragraph']:
 style=doc.styles[style_name]
 style.paragraph_format.line_spacing=1.5
 style.paragraph_format.space_after=Pt(3)
 if style_name.startswith('Heading'):
  style.paragraph_format.space_before=Pt(8)
  style.paragraph_format.space_after=Pt(3)

for paragraph in list(doc.paragraphs):
 text=paragraph.text
 paragraph.paragraph_format.line_spacing=1.5
 paragraph.paragraph_format.page_break_before=False
 paragraph.paragraph_format.widow_control=True
 if paragraph.style.name.startswith('Heading'):
  paragraph.paragraph_format.keep_with_next=True
  if text.endswith('：'):
   paragraph.text=text[:-1]
  continue
 if paragraph.style.name=='Title':continue
 is_bullet=text.startswith('• ')
 clean_text=text[2:] if is_bullet else text
 match=re.match(r'^([^：]{1,30})：\s*(.+)$',clean_text)
 if not match:
  match=re.match(r'^(公共能力抽象|可视化适配层|配置驱动链路|缓存与交付)\s{2,}(.+)$',clean_text)
 if not match:continue
 title_text,body_text=match.groups()
 body_element=OxmlElement('w:p')
 paragraph._p.addnext(body_element)
 body_paragraph=Paragraph(body_element,paragraph._parent)
 body_paragraph.style='Normal'
 body_paragraph.add_run(body_text)
 body_paragraph.paragraph_format.line_spacing=1.5
 body_paragraph.paragraph_format.space_after=Pt(4)
 body_paragraph.paragraph_format.widow_control=True
 if is_bullet:
  body_paragraph.paragraph_format.left_indent=Cm(.32)
 paragraph.clear()
 paragraph.style='Heading 3'
 paragraph.add_run(('• ' if is_bullet else '')+title_text)
 paragraph.paragraph_format.left_indent=Cm(.32) if is_bullet else Cm(0)
 paragraph.paragraph_format.first_line_indent=Cm(-.32) if is_bullet else Cm(0)
 paragraph.paragraph_format.space_before=Pt(6)
 paragraph.paragraph_format.space_after=Pt(0)
 paragraph.paragraph_format.keep_with_next=True
 paragraph.paragraph_format.line_spacing=1.5

doc.save(r'C:\AIApp\ai-agent\resume-format\刘谦-简历-1.5倍行距版.docx')
print('Updated')

