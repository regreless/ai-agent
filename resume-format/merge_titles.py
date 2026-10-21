from docx import Document
from docx.shared import RGBColor,Pt

source_path=r'C:\AIApp\ai-agent\resume-format\刘谦-简历-1.5倍行距版.docx'
output_path=r'C:\AIApp\ai-agent\resume-format\刘谦-简历-标题同行版.docx'
TITLE_COLOR='1F4E79'
doc=Document(source_path)
paragraphs=list(doc.paragraphs)
merged_count=0
for paragraph_index,paragraph in enumerate(paragraphs[:-1]):
 if paragraph.style.name!='Heading 3':continue
 next_paragraph=paragraphs[paragraph_index+1]
 if next_paragraph.style.name!='Normal' or next_paragraph.text.startswith('• '):continue
 title_text=paragraph.text
 body_text=next_paragraph.text
 paragraph.clear()
 paragraph.style='Normal'
 paragraph.paragraph_format.keep_with_next=False
 paragraph.paragraph_format.keep_together=False
 paragraph.paragraph_format.line_spacing=1.5
 paragraph.paragraph_format.space_before=Pt(3)
 paragraph.paragraph_format.space_after=Pt(2)
 if title_text.startswith('• '):
  paragraph.add_run('• ')
  title_text=title_text[2:]
 title_run=paragraph.add_run(title_text+'：')
 title_run.bold=True
 title_run.font.color.rgb=RGBColor.from_string(TITLE_COLOR)
 body_run=paragraph.add_run(body_text)
 body_run.bold=False
 body_run.font.color.rgb=RGBColor(0,0,0)
 next_paragraph._element.getparent().remove(next_paragraph._element)
 merged_count+=1
for style_name in ['Heading 1','Heading 2','Heading 3']:
 doc.styles[style_name].font.color.rgb=RGBColor.from_string(TITLE_COLOR)
for paragraph in doc.paragraphs:
 if paragraph.style.name.startswith('Heading'):
  for run in paragraph.runs:run.font.color.rgb=RGBColor.from_string(TITLE_COLOR)
doc.save(output_path)
print('Merged labels:',merged_count)
