from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor,Pt

BASE_PATH=Path(r'C:\AIApp\ai-agent\resume-format')
OUTPUT_PATH=BASE_PATH/'刘谦-简历-导航版.docx'
doc=Document(BASE_PATH/'刘谦-简历-微软雅黑版.docx')
STACKS={
 '蓝想数科':'React、NestJS / FastAPI、Prisma、MySQL、LangChain、LangGraph、DeepAgent、RAG、pgvector、DeepSeek、Claude Code、Cursor、MCP、Skills、Rules',
 '上海领鸟云':'React、TypeScript、NestJS、Prisma、MySQL、Docker、TanStack Query、Zustand、Ant Design、ECharts、Vite、Jenkins'
}
in_projects=False
project_name=None
navigation_labels=[]
for paragraph in list(doc.paragraphs):
 if paragraph.text=='项目经验':
  in_projects=True
  continue
 if not in_projects:continue
 if paragraph.style.name=='Heading 2':
  project_name=next((key for key in STACKS if key in paragraph.text),None)
 if not paragraph.text.startswith(('项目背景：','技术栈：')):continue
 label_text,body_text=paragraph.text.split('：',1)
 if label_text=='技术栈':body_text=STACKS[project_name]
 body_element=OxmlElement('w:p')
 paragraph._p.addnext(body_element)
 body_paragraph=Paragraph(body_element,paragraph._parent)
 body_paragraph.style='Normal'
 body_paragraph.paragraph_format.line_spacing=1.5
 body_paragraph.paragraph_format.space_before=Pt(0)
 body_paragraph.paragraph_format.space_after=Pt(0)
 body_paragraph.paragraph_format.keep_with_next=False
 body_run=body_paragraph.add_run(body_text)
 body_run.bold=False
 body_run.font.name='微软雅黑'
 body_run.font.color.rgb=RGBColor.from_string('202020')
 body_run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
 paragraph.clear()
 paragraph.style='Heading 3'
 paragraph.paragraph_format.line_spacing=1.5
 paragraph.paragraph_format.space_before=Pt(0)
 paragraph.paragraph_format.space_after=Pt(0)
 paragraph.add_run(label_text+'：')
 navigation_labels.append((project_name,label_text))
assert len(navigation_labels)==4,navigation_labels
doc.save(OUTPUT_PATH)
print('Prepared four navigation headings')
