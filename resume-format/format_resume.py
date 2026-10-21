from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

source_path=Path(r'C:\Users\Administrator\Desktop\刘谦-简历-精简版.docx')
output_path=Path(r'C:\AIApp\ai-agent\resume-format\刘谦-简历-排版优化版.docx')
source_doc=Document(source_path)
output_doc=Document()
section=output_doc.sections[0]
section.page_width=Cm(21)
section.page_height=Cm(29.7)
section.top_margin=Cm(1.5)
section.bottom_margin=Cm(1.5)
section.left_margin=Cm(1.65)
section.right_margin=Cm(1.65)
section.footer_distance=Cm(.65)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Paragraph']:
 style=output_doc.styles[style_name]
 style.font.name='Calibri'
 style.font.size=Pt(10.5)
 style.font.color.rgb=RGBColor(0,0,0)
 style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
 style.paragraph_format.line_spacing=Pt(13)
 style.paragraph_format.space_after=Pt(2)
 style.paragraph_format.space_before=Pt(0)
for style_name,size in [('Heading 1',13),('Heading 2',11),('Heading 3',10.5)]:
 style=output_doc.styles[style_name]
 style.font.size=Pt(size)
 style.font.bold=True
 style.paragraph_format.space_before=Pt(7 if style_name=='Heading 1' else 4)
 style.paragraph_format.space_after=Pt(3)
 style.paragraph_format.keep_with_next=True

main_indices={4,6,28,58}
sub_indices={7,12,17,24,29,40,50,59,70}
minor_indices={32,35,43,45,60,64,78}
bullet_indices=set(range(8,12))|set(range(13,17))|set(range(18,23))|set(range(25,28))|set(range(36,40))|set(range(46,50))|set(range(53,57))|{61,62,65,66,67}|set(range(73,77))|set(range(79,83))
for index,old_paragraph in enumerate(source_doc.paragraphs):
 text=old_paragraph.text.strip()
 if not text:continue
 if index in main_indices:style_name='Heading 1'
 elif index in sub_indices:style_name='Heading 2'
 elif index in minor_indices:style_name='Heading 3'
 else:style_name='Normal'
 paragraph=output_doc.add_paragraph(style=style_name)
 if index in {28,58}:paragraph.paragraph_format.page_break_before=True
 if index==0:
  paragraph.style='Title'
  paragraph.paragraph_format.space_after=Pt(5)
  run=paragraph.add_run(text)
  run.font.size=Pt(24)
  run.bold=True
 elif index==1:
  run=paragraph.add_run(text)
  run.bold=True
  run.font.size=Pt(12)
  paragraph.paragraph_format.space_after=Pt(5)
 elif index==2:
  paragraph.add_run('6 年工作经验  ｜  17317619053  ｜  tunshisky@outlook.com  ｜  本科  ｜  32岁').font.size=Pt(9.5)
  paragraph.paragraph_format.space_after=Pt(5)
 elif index==3:
  paragraph.add_run(text).font.size=Pt(9.5)
 elif index in {29,40,50}:
  left_text,right_text=text.split('\t')
  paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17.7),WD_TAB_ALIGNMENT.RIGHT)
  paragraph.add_run(left_text.strip()).bold=True
  paragraph.add_run('\t'+right_text.strip()).font.size=Pt(10)
 elif index in bullet_indices:
  import re
  text=re.sub(r'^\d+\.\s*','',text)
  paragraph.paragraph_format.left_indent=Cm(.32)
  paragraph.paragraph_format.first_line_indent=Cm(-.32)
  paragraph.add_run('• ')
  if '：' in text and index not in {15,16,18,19,20,36,37,38,39,46,47,48,49,53,54,55,56,61,62}:
   lead_text,body_text=text.split('：',1)
   if len(lead_text)<32:
    paragraph.add_run(lead_text+'：').bold=True
    paragraph.add_run(body_text)
   else:paragraph.add_run(text)
  else:paragraph.add_run(text)
 else:
  if index in {30,31,33,34,41,42,44,51,52,63,69,71,72,77,83}:
   import re
   match=re.match(r'^(角色定位|技术栈|项目概述|项目价值|项目背景|交付结果)([：\s]+)(.*)$',text)
   if match:
    paragraph.add_run(match.group(1)+'：').bold=True
    paragraph.add_run(match.group(3))
   else:paragraph.add_run(text)
  else:paragraph.add_run(text)
 paragraph.paragraph_format.widow_control=True
footer_paragraph=section.footer.paragraphs[0]
footer_paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer_paragraph.add_run()
run.font.size=Pt(8)
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
run._r.addnext(field)
output_doc.core_properties.author='刘谦'
output_doc.core_properties.title='刘谦 简历'
output_doc.core_properties.subject=''
output_doc.save(output_path)
print(output_path)
