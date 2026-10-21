from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

SOURCE_PATH = Path('C:/Users/Administrator/Desktop/inter.docx')
OUTPUT_PATH = Path('C:/Users/Administrator/Desktop/inter_重新排版.docx')
source_doc = Document(SOURCE_PATH)
main_text = source_doc.paragraphs[1].text
main_text = re.sub(r'第\s*[23]\s*页', '', main_text)
main_text = re.sub(r'(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff，。；：、）])', '', main_text)
main_text = re.sub(r'\s+([，。；：、）])', r'\1', main_text)
heading_names = ['职业概述', '专业技能', 'React 与 TypeScript', 'ToB 业务、数据平台与工程化', '智能体协作与 AI 工程化', 'Node.js / NestJS 全栈能力', '跨端与团队协作', '工作经历', '蓝想数科｜前端开发（前后端协作与 AI 工程化）', '上海领鸟云｜前端开发', '上海联蔚数科｜前端开发', '项目经验', '企业级数据智能与低代码分析平台｜配置驱动的可视化体系', '急诊重症临床协同平台｜高可靠流程、可观测性与智能体工程']
heading_pattern = '|'.join(re.escape(v) + (r'(?![，。])' if v == '项目经验' else '') for v in sorted(heading_names, key=len, reverse=True))
split_pattern = '(' + heading_pattern + r'|(?<!\S)[1-6]\.\s+|角色定位：|前端技术：|后端技术：|AI 工程：|主要技术：|项目背景：|架构设计与关键实践：|项目价值：|工作经验：)'
parts = re.split(split_pattern, main_text)
blocks = []
pending_label = ''
for part in parts:
    part = part.strip()
    if not part:
        continue
    if part in heading_names:
        blocks.append(('heading', part))
    elif re.fullmatch(r'[1-6]\.', part):
        pending_label = part + ' '
    elif part == '架构设计与关键实践：':
        blocks.append(('label', part))
    elif part.endswith('：'):
        pending_label = part
    else:
        blocks.append(('body', pending_label + part))
        pending_label = ''

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
section.footer_distance = Cm(.65)

def set_font(font_obj, font_size, bold=False):
    font_obj.name = 'Calibri'
    font_obj.size = Pt(font_size)
    font_obj.bold = bold
    font_obj.color.rgb = RGBColor(0, 0, 0)

for style_name, font_size, bold in [('Normal',10.5,False),('Title',24,True),('Heading 1',14,True),('Heading 2',11.5,True)]:
    style = doc.styles[style_name]
    set_font(style.font, font_size, bold)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    paragraph_format = style.paragraph_format
    paragraph_format.space_after = Pt(4)
    paragraph_format.line_spacing = Pt(14)
    paragraph_format.widow_control = True
    if style_name.startswith('Heading'):
        paragraph_format.space_before = Pt(9 if style_name == 'Heading 1' else 6)
        paragraph_format.space_after = Pt(5)
        paragraph_format.keep_with_next = True

for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def add_body(text, compact=False):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    if re.match(r'^\d\.', text):
        paragraph.paragraph_format.left_indent = Cm(.45)
        paragraph.paragraph_format.first_line_indent = Cm(-.45)
    elif '：' in text and text.index('：') < 12:
        label, value = text.split('：', 1)
        paragraph.clear()
        paragraph.add_run(label + '：').bold = True
        paragraph.add_run(value)
    return paragraph

doc.add_paragraph('刘谦', 'Title')
job_title = source_doc.paragraphs[0].text.split('求职岗位：',1)[1].strip()
paragraph = doc.add_paragraph('求职岗位：' + job_title)
paragraph.runs[0].bold = True
paragraph.paragraph_format.space_after = Pt(7)

for block_kind, text in blocks:
    if text.startswith('核心能力：'):
        add_body(text, True)
    elif text.startswith('工作经验：'):
        add_body(text)
    elif block_kind == 'heading':
        heading_level = 1 if text in ['职业概述','专业技能','工作经历','项目经验'] else 2
        paragraph = doc.add_paragraph(text, 'Heading ' + str(heading_level))
        if text in ['Node.js / NestJS 全栈能力', '上海联蔚数科｜前端开发', '急诊重症临床协同平台｜高可靠流程、可观测性与智能体工程']:
            paragraph.paragraph_format.page_break_before = True
    elif block_kind == 'label':
        paragraph = add_body(text)
        paragraph.runs[0].bold = True
        paragraph.paragraph_format.keep_with_next = True
    else:
        # The overview contains two distinct paragraphs in the source text.
        if text.startswith('6 年以上') and '前端是主要优势' in text:
            first_part, second_part = text.split('前端是主要优势',1)
            add_body(first_part.strip())
            add_body('前端是主要优势' + second_part)
        else:
            add_body(text)

project_title, project_stack = source_doc.paragraphs[5].text.split('\n', 1)
doc.add_paragraph(project_title, 'Heading 2')
add_body(project_stack)
add_body(source_doc.paragraphs[6].text)
for item_index in range(7, 11):
    add_body(str(item_index - 6) + '. ' + source_doc.paragraphs[item_index].text)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
set_font(run.font, 9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)
doc.core_properties.title = '刘谦 求职简历'
doc.core_properties.subject = '中高级开发工程师 全栈偏前端'
doc.save(OUTPUT_PATH)
print(str(OUTPUT_PATH).encode('unicode_escape').decode())
print('blocks', len(blocks), 'characters', sum(len(p.text) for p in doc.paragraphs))
