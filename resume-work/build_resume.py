from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT

OUTPUT_PATH = r'C:\Users\Administrator\Desktop\刘谦-简历-结构优化版.docx'
doc = Document()
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)
for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[style_name]
    style.font.name = '微软雅黑'
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(4)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(2)
for style_name, font_size in [('Heading 1', 13), ('Heading 2', 10.5)]:
    style = doc.styles[style_name]
    style.font.size = Pt(font_size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(11 if style_name == 'Heading 1' else 7)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True

def add_text(text, bold_label=None, size=None):
    paragraph = doc.add_paragraph()
    if bold_label:
        paragraph.add_run(bold_label).bold = True
        paragraph.add_run(text)
    else:
        paragraph.add_run(text)
    if size:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    paragraph.paragraph_format.widow_control = True
    return paragraph

def add_item(label, text):
    paragraph = add_text(text, label + '  ')
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph

def add_role(company, role, date):
    paragraph = doc.add_paragraph(style='Heading 2')
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(company + '  |  ' + role)
    paragraph.add_run('\t' + date).font.size = Pt(9.5)

doc.add_paragraph('刘谦', 'Title')
paragraph = add_text('AI 应用开发 / AI 全栈', size=12)
paragraph.paragraph_format.space_after = Pt(5)
add_text('17317619053  |  tunshisky@outlook.com  |  本科  |  32 岁  |  6 年工作经验', size=9)

doc.add_paragraph('职业概述', 'Heading 1')
add_text('具备 SaaS 数据平台、医疗业务系统及多端企业应用开发经验，能够独立推进需求澄清、方案设计、开发联调、测试发布与线上排查。以前端架构和复杂业务交付为基础，具备 Node.js 全栈及 RAG、Agent、流式交互实践，能够贯通页面、接口、数据与 AI 能力集成。')

source_doc = Document(r'C:\Users\Administrator\Desktop\刘谦-简历-精简版.docx')
doc.add_paragraph('专业技能', 'Heading 1')
for source_index in range(9, 30):
    source_text = source_doc.paragraphs[source_index].text.strip()
    if not source_text:
        continue
    if source_index == 19:
        doc.add_page_break()
        doc.add_paragraph('专业技能 续', 'Heading 1')
    if source_index in [9, 14, 19, 26]:
        doc.add_paragraph(source_text, 'Heading 2')
    else:
        add_text(source_text)

doc.add_paragraph('工作经历', 'Heading 1')
add_role('蓝想数科', 'AI 应用', '2025.06 — 至今')
add_text('负责预检分诊、急诊抢救、重症监护、麻醉随访等模块的需求落地与完整交付，协调前后端联调、测试发布及线上问题排查。')
add_text('参与接口契约、监控数据及智能体研发规范建设，推动 AI 辅助研发流程在业务交付中应用。')

add_role('上海领鸟云', '前端开发', '2022.07 — 2025.06')
add_text('负责客户管理、数据分析、选址分析、产业图谱等 SaaS 模块及低代码配置平台的核心开发，承担业务迭代与平台公共能力建设。')

add_role('上海联蔚数科', '前端开发', '2019.09 — 2022.07')
add_text('负责微信小程序、H5 与 PC 管理后台的模块开发、接口联调及上线交付；使用 Redux 处理跨页面状态与持久化，分层封装授权、分享和平台 API 差异，并实施路由懒加载、请求合并与分包优化。')

doc.add_page_break()
doc.add_paragraph('项目经验', 'Heading 1')
doc.add_paragraph('急诊重症临床协同平台', 'Heading 2')
add_text('所属公司  蓝想数科', size=9)
add_item('项目背景', '面向时间敏感、多角色协同及数据准确性要求高的急诊重症场景，将分散的诊疗数据组织为连续业务流程。')
add_item('技术栈', 'React、TypeScript、Node.js、NestJS、Prisma、MySQL、Playwright。')
add_item('统一数据模型', '按指标值、单位、采集时间和状态组织生命体征及诊疗数据，集中处理异常标记、数据缺失与接口失败，减少不同页面对同类数据的重复解释。')
add_item('状态与规则分层', '以抢救阶段、处置节点和关键时间表达业务进度及异常状态，保留操作追溯线索；将表单拆为字段配置、字典数据、校验规则、动态联动和提交转换五层，用校验矩阵衔接需求、开发与测试。')
add_item('质量与可观测性', '结合核心流程自动化回归、类型检查与生产构建，覆盖正常流程、编辑回显和状态流转；将监控字段与后端错误码对齐，打通用户操作、接口异常到问题定位的链路。')
add_item('交付结果', '形成结构化规则、分层实现、监控反馈、自动化验证和人工验收的交付流程，为高频迭代提供回归与问题追踪依据。')

doc.add_paragraph('企业级数据智能与低代码分析平台', 'Heading 2')
add_text('所属公司  上海领鸟云', size=9)
add_item('项目背景', '面向多数据源、多分析维度及频繁变化的企业分析需求，支持数据查询、空间分析、关系洞察和自助式页面搭建。')
add_item('技术栈', 'React、TypeScript、TanStack Query、Zustand、Ant Design、ECharts、Vite、Jenkins。')
add_item('公共能力抽象', '将查询、数据展示、字典解析、交互操作与详情编辑拆为可组合单元，通过输入输出契约明确复用边界，支持业务模块按需组装。')
add_item('可视化适配层', '将后端差异化数据转换为统一图表模型并生成 ECharts 配置，集中处理趋势、分类、地图及关系数据的联动、缩放、自适应与异常兜底。')
add_item('配置驱动链路', '描述页面结构、组件属性和数据源关系，贯通属性编辑、配置保存、页面预览及运行时渲染；通过组件注册机制支持后续扩展。')
add_item('缓存与交付', '使用 TanStack Query 管理服务端数据缓存与失效，通过 Vite 构建拆分和 Jenkins 流水线支持检查、构建及多环境交付。')
add_item('交付结果', '沉淀统一的数据接入、图表表达与页面搭建模型，将分散的业务页面开发转化为可复用的平台能力。')

doc.core_properties.title = '刘谦个人简历'
doc.core_properties.subject = 'AI应用开发与全栈岗位简历'
doc.core_properties.author = '刘谦'
doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
